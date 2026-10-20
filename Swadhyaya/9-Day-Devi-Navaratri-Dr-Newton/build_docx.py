"""Build one Word file from the Devi Navaratri HTML readers.

Fonts and sizes follow the pages:

- day titles: Libre Baskerville, 19px
- section titles: Libre Baskerville, 18px
- reading text: Noto Sans Telugu, 14px, line-height 1.73
- programme page: Georgia, at the sizes used in index.html

Run from this folder:

    python build_docx.py
"""

from __future__ import annotations

import json
import re
import uuid
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
FONT_DIR = ROOT / "assets" / "fonts"
BANNER = ROOT / "assets" / "devi-navaratrula-jnanam-banner.png"
OUTPUT = ROOT / "Devi-Navaratrula-Jnanam.docx"

# CSS pixels at 96dpi are 0.75 points.
def px(value: float) -> Pt:
    return Pt(value * 0.75)


HEADING = RGBColor(0x6B, 0x4A, 0x25)
INDEX_HEADING = RGBColor(0x6E, 0x4A, 0x25)
INK = RGBColor(0x35, 0x2D, 0x26)
MUTED = RGBColor(0x56, 0x48, 0x3A)
SOFT = RGBColor(0x8A, 0x6A, 0x4A)
GOLD = RGBColor(0xA1, 0x7B, 0x43)
ACCENT = RGBColor(0xBA, 0x7A, 0x3A)
SUB = RGBColor(0x7C, 0x6E, 0x5D)

SERIF = "Libre Baskerville"
BODY = "Noto Sans Telugu"
UI = "Source Sans 3"
INDEX = "Georgia"

FONTS = {
    SERIF: {
        "regular": FONT_DIR / "LibreBaskerville-Regular.ttf",
        "bold": FONT_DIR / "LibreBaskerville-Bold.ttf",
    },
    BODY: {"regular": FONT_DIR / "NotoSansTelugu-Regular.ttf"},
    UI: {"regular": FONT_DIR / "SourceSans3-Regular.ttf"},
}


def load_programme() -> list[dict]:
    html = (ROOT / "index.html").read_text(encoding="utf-8")
    days = []
    for match in re.finditer(
        r'<li id="day-(\d+)">(?:<a href="[^"]*">|<div class="pending">)<strong>(.*?)</strong><em>(.*?)</em>',
        html,
    ):
        days.append(
            {
                "number": int(match.group(1)),
                "label": re.sub(r"<[^>]+>", "", match.group(2)).strip(),
                "summary": re.sub(r"<[^>]+>", "", match.group(3)).strip(),
            }
        )
    return days


def load_chapters(day_number: int) -> list[dict]:
    path = ROOT / f"day-{day_number:02d}-en.html"
    if not path.exists():
        return []
    text = path.read_text(encoding="utf-8")
    marker = "const chapters="
    start = text.find(marker)
    if start < 0:
        return []
    chapters, _ = json.JSONDecoder().raw_decode(text[start + len(marker) :])
    return chapters


def set_run_font(run, name: str, size: Pt, color: RGBColor, bold: bool = False) -> None:
    # Word stores size in half-points, so 19px (14.25pt) becomes 14.5pt.
    half_points = int(size.pt * 2 + 0.5)
    run.bold = bold
    run.italic = False
    run.font.size = Pt(half_points / 2)
    run.font.name = name
    run.font.color.rgb = color
    run.font.underline = False
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        r_fonts.set(qn(attr), name)
    color_el = r_pr.find(qn("w:color"))
    if color_el is not None and qn("w:themeColor") in color_el.attrib:
        del color_el.attrib[qn("w:themeColor")]


def set_line_spacing(paragraph, multiplier: float, after_px: float = 0, before_px: float = 0) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = multiplier
    fmt.space_after = px(after_px)
    fmt.space_before = px(before_px)


class Builder:
    def __init__(self) -> None:
        self.doc = Document()
        self._bookmark_id = 1
        self._prepare_section()
        self._prepare_styles()

    def _prepare_section(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = footer.add_run("Jnanodayam · Devi Navaratrula Jnanam · Dr. Newton Kondaveti · Quantum Life University")
        set_run_font(run, INDEX, px(12), SOFT)
        set_line_spacing(footer, 1.2)

    def _prepare_styles(self) -> None:
        normal = self.doc.styles["Normal"]
        normal.font.name = BODY
        normal.font.size = px(14)
        normal.font.color.rgb = INK
        r_pr = normal.element.get_or_add_rPr()
        r_fonts = r_pr.find(qn("w:rFonts"))
        if r_fonts is None:
            r_fonts = OxmlElement("w:rFonts")
            r_pr.append(r_fonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            r_fonts.set(qn(attr), BODY)

    def bookmark(self, paragraph, name: str) -> None:
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(self._bookmark_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(self._bookmark_id))
        self._bookmark_id += 1
        paragraph._p.insert(0, start)
        paragraph._p.append(end)

    def add_text(self, paragraph, text: str, name: str, size: Pt, color: RGBColor, bold: bool = False):
        run = paragraph.add_run(text)
        set_run_font(run, name, size, color, bold)
        return run

    def add_link(self, paragraph, text: str, anchor: str, name: str, size: Pt, color: RGBColor, bold: bool = False) -> None:
        link = OxmlElement("w:hyperlink")
        link.set(qn("w:anchor"), anchor)
        run = OxmlElement("w:r")
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            r_fonts.set(qn(attr), name)
        r_pr.append(r_fonts)
        if bold:
            bold_el = OxmlElement("w:b")
            r_pr.append(bold_el)
        half_points = str(int(size.pt * 2 + 0.5))
        size_el = OxmlElement("w:sz")
        size_el.set(qn("w:val"), half_points)
        size_cs = OxmlElement("w:szCs")
        size_cs.set(qn("w:val"), half_points)
        color_el = OxmlElement("w:color")
        color_el.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "none")
        r_pr.extend([size_el, size_cs, color_el, underline])
        run.append(r_pr)
        text_el = OxmlElement("w:t")
        text_el.set(qn("xml:space"), "preserve")
        text_el.text = text
        run.append(text_el)
        link.append(run)
        paragraph._p.append(link)

    def cover(self) -> None:
        if BANNER.exists():
            width = self.doc.sections[0].page_width - self.doc.sections[0].left_margin - self.doc.sections[0].right_margin
            self.doc.add_picture(str(BANNER), width=width)
            self.doc.paragraphs[-1].paragraph_format.space_after = px(14)

        brand = self.doc.add_paragraph()
        self.add_text(brand, "Jnanodayam", INDEX, px(15), INDEX_HEADING, bold=True)
        set_line_spacing(brand, 1.2, after_px=2)
        sub = self.doc.add_paragraph()
        self.add_text(sub, "Devi Navaratrula Jnanam · Dr. Newton Kondaveti", INDEX, px(12), SOFT)
        set_line_spacing(sub, 1.3, after_px=16)

        title = self.doc.add_paragraph()
        self.bookmark(title, "programme")
        self.add_text(title, "Programme", INDEX, px(22), INDEX_HEADING, bold=True)
        set_line_spacing(title, 1.2, after_px=8)

        lede = self.doc.add_paragraph()
        self.add_text(
            lede,
            "Nine days of Devi Navaratri knowledge with Brahmarshi Chakravarti Dr. Newton Kondaveti of Quantum Life University. The outer festival is the reminder. The work is the inner sadhana — 11 to 20 October 2026.",
            INDEX,
            px(15),
            MUTED,
        )
        set_line_spacing(lede, 1.6, after_px=14)

        note = self.doc.add_paragraph()
        self.add_text(note, "How to use: ", INDEX, px(14), MUTED, bold=True)
        self.add_text(
            note,
            "Open a day and read it section by section. The English follows the Telugu talk for that morning.",
            INDEX,
            px(14),
            MUTED,
        )
        set_line_spacing(note, 1.6, after_px=8)
        p_pr = note._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), "BA7A3A")
        borders.append(left)
        p_pr.append(borders)
        shade = OxmlElement("w:shd")
        shade.set(qn("w:val"), "clear")
        shade.set(qn("w:color"), "auto")
        shade.set(qn("w:fill"), "FFF8EB")
        p_pr.append(shade)
        note.paragraph_format.left_indent = Cm(0.3)

    def contents(self, days: list[dict], chapters_by_day: dict[int, list[dict]]) -> None:
        label = self.doc.add_paragraph()
        self.add_text(label, "The Nine Days", INDEX, px(14), ACCENT, bold=True)
        set_line_spacing(label, 1.3, before_px=18, after_px=10)

        for day in days:
            number = day["number"]
            anchor = f"day-{number:02d}"
            chapters = chapters_by_day.get(number, [])
            head = self.doc.add_paragraph()
            if chapters:
                self.add_link(head, day["label"], anchor, INDEX, px(14), INDEX_HEADING, bold=True)
            else:
                self.add_text(head, day["label"], INDEX, px(14), SOFT, bold=True)
            set_line_spacing(head, 1.3, before_px=8, after_px=0)

            summary = self.doc.add_paragraph()
            self.add_text(summary, day["summary"], INDEX, px(13), MUTED if chapters else SOFT)
            set_line_spacing(summary, 1.4, after_px=2)
            summary.paragraph_format.left_indent = Cm(0.4)

            for index, chapter in enumerate(chapters, start=1):
                item = self.doc.add_paragraph()
                self.add_link(
                    item,
                    chapter["title"],
                    f"{anchor}-s{index:02d}",
                    UI,
                    px(12),
                    INK,
                )
                set_line_spacing(item, 1.35, after_px=1)
                item.paragraph_format.left_indent = Cm(0.9)

    def days(self, days: list[dict], chapters_by_day: dict[int, list[dict]]) -> None:
        for day in days:
            number = day["number"]
            chapters = chapters_by_day.get(number, [])
            heading = self.doc.add_paragraph()
            heading.paragraph_format.page_break_before = True
            self.bookmark(heading, f"day-{number:02d}")
            self.add_text(heading, day["label"], SERIF, px(19), HEADING, bold=True)
            set_line_spacing(heading, 1.25, after_px=2)
            outline = OxmlElement("w:outlineLvl")
            outline.set(qn("w:val"), "0")
            heading._p.get_or_add_pPr().append(outline)

            sub = self.doc.add_paragraph()
            if chapters:
                self.add_text(sub, f"Devi Navaratri · Day {number} (English)", UI, px(11), SUB)
            else:
                self.add_text(sub, day["summary"], UI, px(11), SUB)
            set_line_spacing(sub, 1.3, after_px=4)

            blurb = self.doc.add_paragraph()
            self.add_text(blurb, day["summary"], UI, px(13), MUTED)
            set_line_spacing(blurb, 1.45, after_px=12)

            if not chapters:
                pending = self.doc.add_paragraph()
                self.add_text(pending, "This day’s reading is not in the programme yet.", BODY, px(14), INK)
                set_line_spacing(pending, 1.73, after_px=8)
                continue

            for index, chapter in enumerate(chapters, start=1):
                title = self.doc.add_paragraph()
                self.bookmark(title, f"day-{number:02d}-s{index:02d}")
                self.add_text(title, chapter["title"], SERIF, px(18), HEADING, bold=True)
                set_line_spacing(title, 1.6, before_px=16, after_px=2)
                outline = OxmlElement("w:outlineLvl")
                outline.set(qn("w:val"), "1")
                title._p.get_or_add_pPr().append(outline)

                minutes = chapter.get("minutes")
                if minutes:
                    meta = self.doc.add_paragraph()
                    self.add_text(meta, f"~{minutes} min", UI, px(11), GOLD)
                    set_line_spacing(meta, 1.3, after_px=9)

                for paragraph_text in chapter.get("paras", []):
                    body = self.doc.add_paragraph()
                    self.add_text(body, paragraph_text, BODY, px(14), INK)
                    set_line_spacing(body, 1.73, after_px=8)

    def save(self) -> None:
        self.doc.save(OUTPUT)
        embed_fonts(OUTPUT)


def guid_key(guid: str) -> bytes:
    parts = guid.strip("{}").split("-")
    raw = bytearray()
    for index, part in enumerate(parts):
        chunk = bytes.fromhex(part)
        raw.extend(reversed(chunk) if index < 3 else chunk)
    return bytes(raw)


def obfuscate(data: bytes, guid: str) -> bytes:
    key = guid_key(guid)
    out = bytearray(data)
    for index in range(min(32, len(out))):
        out[index] ^= key[index % 16]
    return bytes(out)


def embed_fonts(path: Path) -> None:
    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    REL = "http://schemas.openxmlformats.org/package/2006/relationships"
    CT = "http://schemas.openxmlformats.org/package/2006/content-types"
    font_rel = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font"

    with zipfile.ZipFile(path, "r") as archive:
        files = {name: archive.read(name) for name in archive.namelist()}

    font_table = etree_from(files["word/fontTable.xml"])
    rels_name = "word/_rels/fontTable.xml.rels"
    if rels_name in files:
        rels = etree_from(files[rels_name])
    else:
        rels = etree_element(f'<Relationships xmlns="{REL}"/>')
    content_types = etree_from(files["[Content_Types].xml"])
    if content_types.find(f"{{{CT}}}Default[@Extension='odttf']") is None:
        default = content_types.makeelement(f"{{{CT}}}Default")
        default.set("Extension", "odttf")
        default.set("ContentType", "application/vnd.openxmlformats-officedocument.obfuscatedFont")
        content_types.insert(0, default)

    existing_ids = {rel.get("Id") for rel in rels}
    next_id = 1

    def new_id() -> str:
        nonlocal next_id
        while f"rIdFont{next_id}" in existing_ids:
            next_id += 1
        value = f"rIdFont{next_id}"
        existing_ids.add(value)
        next_id += 1
        return value

    fonts_root = font_table
    for family, faces in FONTS.items():
        node = None
        for candidate in fonts_root:
            if candidate.tag == f"{{{W}}}font" and candidate.get(qn("w:name")) == family:
                node = candidate
                break
        if node is None:
            node = OxmlElement("w:font")
            node.set(qn("w:name"), family)
            fonts_root.append(node)
        for kind, font_path in faces.items():
            guid = str(uuid.uuid4()).upper()
            rel_id = new_id()
            target = f"fonts/{family.replace(' ', '')}-{kind}.odttf"
            tag = "w:embedBold" if kind == "bold" else "w:embedRegular"
            embed = OxmlElement(tag)
            embed.set(qn("r:id"), rel_id)
            embed.set(qn("w:fontKey"), "{" + guid + "}")
            node.append(embed)
            rel = rels.makeelement(f"{{{REL}}}Relationship")
            rel.set("Id", rel_id)
            rel.set("Type", font_rel)
            rel.set("Target", target)
            rels.append(rel)
            files["word/" + target] = obfuscate(font_path.read_bytes(), guid)

    files["word/fontTable.xml"] = etree_bytes(font_table)
    files[rels_name] = etree_bytes(rels)
    files["[Content_Types].xml"] = etree_bytes(content_types)

    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)


def etree_from(data: bytes):
    from lxml import etree

    return etree.fromstring(data)


def etree_element(xml: str):
    from lxml import etree

    return etree.fromstring(xml.encode("utf-8"))


def etree_bytes(node) -> bytes:
    from lxml import etree

    return etree.tostring(node, xml_declaration=True, encoding="UTF-8", standalone=True)


def main() -> None:
    days = load_programme()
    chapters_by_day = {day["number"]: load_chapters(day["number"]) for day in days}
    builder = Builder()
    builder.cover()
    builder.contents(days, chapters_by_day)
    builder.days(days, chapters_by_day)
    builder.save()
    chapters = sum(len(items) for items in chapters_by_day.values())
    print(f"Wrote {OUTPUT.name} · {len(days)} days · {chapters} sections")


if __name__ == "__main__":
    main()
